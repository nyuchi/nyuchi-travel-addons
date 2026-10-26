"""Tests for trips_xlsx.py. Run with: python3 -m unittest discover -s tools
"""

import csv
import datetime
import io
import os
import tempfile
import unittest

from openpyxl import load_workbook

import trips_xlsx

BOM = "﻿"

TRIPS = (
    BOM
    + "id,ref,modified_gmt,slug,status,title,excerpt,content,meta:wp_travel_trip_duration,meta:wp_travel_lat,terms:activity\n"
    + '3111,,2026-10-26 12:00:00,tz-trip,draft,"\'=SUM(A1)",,"<p>Line one,\nline ""two""</p>",007,-1.29,game-drives|ballooning\n'
    + ",gor,,gorillas,,Gorilla Trek,,,5,,\n"
)

DAYS = (
    BOM
    + "trip_id,trip_ref,day,label,title,date,time,desc\n"
    + '3111,,1,Day 1,Arrive,2026-11-02,,"Land, then transfer"\n'
    + ",gor,1,Day 1,Trek,,07:00,Early start\n"
)

TERMS = (
    BOM
    + "taxonomy,id,slug,parent_slug,status,etag,name,description,field:country,field:best_months,field:difficulty,field:permit_required\n"
    + "travel_locations,10,tanzania,,live,abc123,Tanzania,<p>Old</p>,Tanzania,6|7|8,,\n"
    + "activity,20,ballooning,,live,def456,Ballooning,,,,easy,yes\n"
)


def rows_of(text):
    return list(csv.reader(io.StringIO(text.lstrip(BOM), newline="")))


class ConverterTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.dir.cleanup)
        self.paths = {}

        for name, text in (("trips.csv", TRIPS), ("days.csv", DAYS), ("terms.csv", TERMS)):
            path = os.path.join(self.dir.name, name)
            with open(path, "w", encoding="utf-8", newline="") as handle:
                handle.write(text)
            self.paths[name] = path

        self.xlsx = os.path.join(self.dir.name, "edit.xlsx")
        trips_xlsx.to_xlsx(
            {"Trips": self.paths["trips.csv"], "Days": self.paths["days.csv"], "Terms": self.paths["terms.csv"]},
            self.xlsx,
        )

    def round_trip(self, name):
        out = os.path.join(self.dir.name, "out")
        trips_xlsx.to_csv(self.xlsx, out)
        with open(os.path.join(out, name), encoding="utf-8-sig", newline="") as handle:
            return list(csv.reader(handle))

    def test_round_trip_is_lossless(self):
        for name, text in (("trips.csv", TRIPS), ("days.csv", DAYS), ("terms.csv", TERMS)):
            with self.subTest(name):
                self.assertEqual(self.round_trip(name), rows_of(text))

    def test_formula_like_and_leading_zero_cells_stay_text(self):
        sheet = load_workbook(self.xlsx)["Trips"]
        title = sheet["F2"]
        self.assertEqual(title.value, "'=SUM(A1)")
        self.assertEqual(title.data_type, "s")
        self.assertEqual(sheet["I2"].value, "007")
        self.assertEqual(sheet["I2"].number_format, "@")

    def test_unguarded_leading_equals_is_never_a_formula(self):
        path = os.path.join(self.dir.name, "raw.csv")
        with open(path, "w", encoding="utf-8", newline="") as handle:
            handle.write("id,title\n1,=HYPERLINK(\"http://x\")\n")
        out = os.path.join(self.dir.name, "raw.xlsx")
        trips_xlsx.to_xlsx({"Trips": path}, out)
        self.assertEqual(load_workbook(out)["Trips"]["B2"].data_type, "s")

    def test_formatting(self):
        workbook = load_workbook(self.xlsx)
        self.assertEqual(workbook.sheetnames, ["Read me", "Trips", "Days", "Terms"])

        trips = workbook["Trips"]
        self.assertEqual(trips.freeze_panes, "C2")
        self.assertEqual(trips["A1"].font.bold, True)
        self.assertTrue(trips.auto_filter.ref)
        self.assertEqual(trips["A2"].fill.fgColor.rgb[-6:], "E7E6E6")  # id is reference only
        self.assertNotEqual(trips["F2"].fill.fgColor.rgb[-6:], "E7E6E6")  # title is editable
        self.assertEqual(trips.column_dimensions["H"].width, trips_xlsx.WIDE_COLUMNS["content"])

        terms = workbook["Terms"]
        lists = {dv.formula1 for dv in terms.data_validations.dataValidation}
        self.assertIn('"easy,moderate,strenuous"', lists)
        self.assertIn('"yes,no"', lists)

    def test_edited_workbook_comes_back_as_csv(self):
        workbook = load_workbook(self.xlsx)
        workbook["Trips"]["F2"] = "Tanzania Deluxe"
        workbook["Terms"]["I3"] = "[clear]"
        workbook.save(self.xlsx)

        trips = self.round_trip("trips.csv")
        self.assertEqual(trips[1][5], "Tanzania Deluxe")
        self.assertEqual(self.round_trip("terms.csv")[2][8], "[clear]")

    def test_numbers_and_dates_typed_in_excel_come_back_as_text(self):
        workbook = load_workbook(self.xlsx)
        trips = workbook["Trips"]
        trips["I2"] = 7.0
        workbook["Days"]["F2"] = datetime.datetime(2026, 11, 3)
        workbook.save(self.xlsx)

        self.assertEqual(self.round_trip("trips.csv")[1][8], "7")
        self.assertEqual(self.round_trip("days.csv")[1][5], "2026-11-03")

    def test_blank_rows_and_columns_from_formatting_are_dropped(self):
        workbook = load_workbook(self.xlsx)
        workbook["Trips"]["Z40"] = None
        workbook["Trips"].cell(row=30, column=3).value = ""
        workbook.save(self.xlsx)

        out = self.round_trip("trips.csv")
        self.assertEqual(len(out), 3)
        self.assertEqual({len(row) for row in out}, {len(out[0])})

    def test_only_the_sheets_given_are_written(self):
        out = os.path.join(self.dir.name, "just-terms.xlsx")
        trips_xlsx.to_xlsx({"Terms": self.paths["terms.csv"]}, out)
        self.assertEqual(load_workbook(out).sheetnames, ["Read me", "Terms"])

    def test_cell_too_long_for_excel_is_refused(self):
        path = os.path.join(self.dir.name, "long.csv")
        with open(path, "w", encoding="utf-8", newline="") as handle:
            handle.write("id,content\n1," + "x" * (trips_xlsx.MAX_CELL + 1) + "\n")
        with self.assertRaises(SystemExit):
            trips_xlsx.to_xlsx({"Trips": path}, os.path.join(self.dir.name, "long.xlsx"))

    def test_semicolon_csv_is_read(self):
        path = os.path.join(self.dir.name, "semi.csv")
        with open(path, "w", encoding="utf-8", newline="") as handle:
            handle.write("id;title\n1;Hello, world\n")
        self.assertEqual(trips_xlsx.read_csv_file(path), [["id", "title"], ["1", "Hello, world"]])

    def test_workbook_without_known_sheets_is_refused(self):
        from openpyxl import Workbook

        path = os.path.join(self.dir.name, "other.xlsx")
        workbook = Workbook()
        workbook.active.title = "Sheet1"
        workbook.save(path)
        with self.assertRaises(SystemExit):
            trips_xlsx.to_csv(path, os.path.join(self.dir.name, "o"))


if __name__ == "__main__":
    unittest.main()
