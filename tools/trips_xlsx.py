#!/usr/bin/env python3
"""Convert the plugin's trip and term CSV files to a formatted Excel workbook, and back.

This is an offline helper. It never touches the website and nothing in the
plugin reads Excel files: the workbook is only a nicer editor for the CSV that
`export-trips-csv` and `export-terms-csv` produce, and `to-csv` turns it back
into the CSV that `import-trips-csv` and `import-terms-csv` accept.

    trips_xlsx.py to-xlsx --trips trips.csv --days days.csv --terms terms.csv -o edit.xlsx
    trips_xlsx.py to-csv edit.xlsx --out-dir ./csv

Rules that carry over from the CSV format (also written into the workbook's
"Read me" sheet): a blank cell leaves a value unchanged, [clear] clears it, grey
columns are reference only and are not applied, and months are numbers joined by |.

Requires: python3 and openpyxl (pip install openpyxl).
"""

import argparse
import csv
import datetime
import io
import os
import sys

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover - exercised only without the dependency
    sys.exit("openpyxl is required: pip install openpyxl")

# Trip content can be long; the csv module's default 128 KB field limit is too low.
csv.field_size_limit(min(sys.maxsize, 2**31 - 1))

# Sheet name -> CSV file name. The order is the tab order.
SHEETS = (("Trips", "trips.csv"), ("Days", "days.csv"), ("Terms", "terms.csv"))
README_SHEET = "Read me"

# Columns the plugin reads back as reference only, or that identify the row.
# They are shaded grey so they are not edited by accident.
REFERENCE_COLUMNS = {
    "Trips": {"id", "modified_gmt", "status"},
    "Days": {"trip_id"},
    "Terms": {"taxonomy", "id", "slug", "parent_slug", "status", "etag"},
}

# Columns that hold long text and should wrap, with a width in characters.
WIDE_COLUMNS = {
    "content": 70,
    "excerpt": 50,
    "desc": 70,
    "description": 60,
    "field:long_description": 60,
    "field:visa_note": 40,
    "field:faqs": 60,
}

# Drop-downs for terms.csv: column -> allowed values. A blank is always allowed.
TERM_LISTS = {
    "field:difficulty": ["easy", "moderate", "strenuous"],
    "field:permit_required": ["yes", "no"],
}

# Excel refuses a cell longer than this; a CSV cell that long cannot round-trip.
MAX_CELL = 32767

READ_ME = (
    "How to edit this workbook",
    "",
    "1. Edit the white cells. Grey columns identify a row or show reference values and are ignored on import.",
    "2. A blank cell means leave the current value unchanged. To clear a value, type [clear].",
    "3. Months are numbers joined by a vertical bar, for example 6|7|8 for June to August.",
    "4. permit_required is yes or no. difficulty is easy, moderate or strenuous.",
    "5. A trip with rows on the Days sheet has its whole day list replaced by those rows.",
    "6. A row on the Trips sheet with no id creates a new draft trip (the import must be run with allow_create).",
    "   Give it a ref, and put the same ref in trip_ref on the Days sheet to attach its days.",
    "7. Terms cannot be created here. Add the term in WordPress first, then export again.",
    "8. Do not change the header row, and do not add or reorder columns you do not recognise.",
    "9. Save as .xlsx, run: trips_xlsx.py to-csv this-file.xlsx --out-dir ./csv, then upload the CSV with the",
    "   import-trips-csv / import-terms-csv abilities. Always do a dry run first.",
)


def read_csv_file(path):
    """Read a CSV file into a list of rows of strings, keeping quoted newlines."""
    with open(path, "r", encoding="utf-8-sig", newline="") as handle:
        text = handle.read()

    first = text.split("\n", 1)[0]
    counts = {",": first.count(","), ";": first.count(";"), "\t": first.count("\t")}
    delimiter = max(counts, key=counts.get)

    rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    return [row for row in rows if any(cell.strip() for cell in row)]


def cell_text(value):
    """A worksheet value as the string the CSV should hold."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    # A typed date can still arrive as a date despite the text format. Dates in
    # the plugin are plain ISO strings, so give back that, without a time of day.
    if isinstance(value, datetime.datetime):
        return value.date().isoformat() if value.time() == datetime.time(0) else value.isoformat(sep=" ")
    if isinstance(value, datetime.date):
        return value.isoformat()
    return str(value)


def write_sheet(workbook, name, rows):
    sheet = workbook.create_sheet(name)
    header = rows[0]
    reference = REFERENCE_COLUMNS.get(name, set())

    grey = PatternFill("solid", fgColor="E7E6E6")
    head_fill = PatternFill("solid", fgColor="1F3A2E")
    head_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")

    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            if len(value) > MAX_CELL:
                sys.exit(
                    f"{name}: row {r}, column '{header[c - 1]}' is {len(value)} characters; "
                    f"Excel cells hold at most {MAX_CELL}. Edit that field in WordPress instead."
                )

            cell = sheet.cell(row=r, column=c, value=value)
            # Everything is text, so Excel cannot turn 007 into 7, 6|7 into a
            # date, or a leading = into a formula.
            cell.data_type = "s"
            cell.number_format = "@"

            column = header[c - 1] if c - 1 < len(header) else ""
            if r == 1:
                cell.fill = head_fill
                cell.font = head_font
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            else:
                cell.alignment = wrap
                if column in reference:
                    cell.fill = grey

    for c, column in enumerate(header, start=1):
        letter = get_column_letter(c)
        if column in WIDE_COLUMNS:
            sheet.column_dimensions[letter].width = WIDE_COLUMNS[column]
        else:
            longest = max((len(row[c - 1]) for row in rows[1:] if c - 1 < len(row)), default=0)
            sheet.column_dimensions[letter].width = min(max(len(column) + 2, longest + 2, 8), 40)

    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "B2" if name == "Days" else "C2"
    sheet.auto_filter.ref = sheet.dimensions

    if name == "Terms":
        for c, column in enumerate(header, start=1):
            if column in TERM_LISTS and len(rows) > 1:
                letter = get_column_letter(c)
                validation = DataValidation(
                    type="list",
                    formula1='"' + ",".join(TERM_LISTS[column]) + '"',
                    allow_blank=True,
                    showErrorMessage=True,
                    errorTitle="Not an allowed value",
                    error="Choose from the list, or leave blank to keep the current value.",
                )
                sheet.add_data_validation(validation)
                validation.add(f"{letter}2:{letter}{len(rows)}")

    return sheet


def to_xlsx(paths, output):
    workbook = Workbook()
    workbook.remove(workbook.active)

    readme = workbook.create_sheet(README_SHEET)
    for r, line in enumerate(READ_ME, start=1):
        cell = readme.cell(row=r, column=1, value=line)
        cell.data_type = "s"
        if r == 1:
            cell.font = Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 120

    written = 0
    for name, _ in SHEETS:
        path = paths.get(name)
        if not path:
            continue

        rows = read_csv_file(path)
        if not rows:
            sys.exit(f"{path} is empty.")

        write_sheet(workbook, name, rows)
        written += 1

    if not written:
        sys.exit("Give at least one of --trips, --days or --terms.")

    workbook.save(output)
    return written


def to_csv(source, out_dir):
    workbook = load_workbook(source, data_only=True, read_only=False)
    os.makedirs(out_dir, exist_ok=True)
    produced = []

    for name, filename in SHEETS:
        if name not in workbook.sheetnames:
            continue

        sheet = workbook[name]
        rows = [[cell_text(cell.value) for cell in row] for row in sheet.iter_rows()]

        # Trailing blank columns are formatting leftovers, not data.
        width = 0
        for row in rows:
            for i, value in enumerate(row, start=1):
                if value != "":
                    width = max(width, i)
        rows = [row[:width] for row in rows if any(cell != "" for cell in row)]

        if not rows:
            continue

        header = rows[0]
        if any(column == "" for column in header):
            sys.exit(f"{name}: the header row has a blank column name. Do not leave gaps in the header.")

        path = os.path.join(out_dir, filename)
        # utf-8-sig writes the BOM Excel expects; the plugin strips it on import.
        with open(path, "w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle, lineterminator="\n")
            for row in rows:
                writer.writerow(row + [""] * (len(header) - len(row)))

        produced.append(path)

    if not produced:
        sys.exit(f"{source} has no Trips, Days or Terms sheet.")

    return produced


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    commands = parser.add_subparsers(dest="command", required=True)

    up = commands.add_parser("to-xlsx", help="build a workbook from the plugin's CSV files")
    up.add_argument("--trips", help="trips.csv from export-trips-csv")
    up.add_argument("--days", help="days.csv from export-trips-csv")
    up.add_argument("--terms", help="terms.csv from export-terms-csv")
    up.add_argument("-o", "--output", required=True, help="workbook to write (.xlsx)")

    down = commands.add_parser("to-csv", help="write the CSV files back out of an edited workbook")
    down.add_argument("workbook", help="the edited .xlsx")
    down.add_argument("--out-dir", required=True, help="folder for trips.csv, days.csv and terms.csv")

    args = parser.parse_args(argv)

    if args.command == "to-xlsx":
        count = to_xlsx({"Trips": args.trips, "Days": args.days, "Terms": args.terms}, args.output)
        print(f"Wrote {args.output} ({count} sheet{'s' if count != 1 else ''} plus Read me).")
    else:
        for path in to_csv(args.workbook, args.out_dir):
            print(f"Wrote {path}")


if __name__ == "__main__":
    main()
